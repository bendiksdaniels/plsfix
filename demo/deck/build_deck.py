"""Build the pls,fix PowerPoint pickers demo deck: a four-slide walkthrough
of the Slide/Where pickers. SLIDE/MARGIN/GAP below must equal
src/link/status.ts SLIDE_16_9 and src/ppt/placement.ts SLIDE_MARGIN/GAP, so
slides 3-4 never drift from the pane's geometry. Python: no Rust PPTX crate."""

from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Pt

# Slide geometry in points: src/link/status.ts SLIDE_16_9, src/ppt/placement.ts
# SLIDE_MARGIN and SLIDE_GAP. The content area is the slide minus the margin on
# every side; halves split that area with the gap between them.
SLIDE_WIDTH = Pt(960)
SLIDE_HEIGHT = Pt(540)
MARGIN = Pt(36)
GAP = Pt(12)

CONTENT_LEFT = MARGIN
CONTENT_TOP = MARGIN
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
CONTENT_HEIGHT = SLIDE_HEIGHT - 2 * MARGIN

HALF_WIDTH = (CONTENT_WIDTH - GAP) // 2  # exact: divides evenly, never a float
LEFT_HALF = (CONTENT_LEFT, CONTENT_TOP, HALF_WIDTH, CONTENT_HEIGHT)
RIGHT_HALF = (CONTENT_LEFT + HALF_WIDTH + GAP, CONTENT_TOP, HALF_WIDTH, CONTENT_HEIGHT)

# The pane's free-space scan (src/ppt/placement.ts occupiedBoxes, src/layout.ts
# placeInFreeSpace) blocks anything whose top is less than a shape's bottom + GAP,
# so the title (text + rule) must end by here for CONTENT_TOP to read as free.
TITLE_BAND_LIMIT = CONTENT_TOP - GAP  # 24 pt
TITLE_HEIGHT = Pt(18)
RULE_HEIGHT = Pt(2)

# Slide 4's caption band: CAPTION_HEIGHT is the one named constant, GAP is
# the breathing room above it, so picture + GAP + caption ties out to
# CONTENT_HEIGHT exactly instead of three independently hand-tuned numbers.
CAPTION_HEIGHT = Pt(40)
PICTURE_HEIGHT = CONTENT_HEIGHT - GAP - CAPTION_HEIGHT
CAPTION_TOP = CONTENT_TOP + CONTENT_HEIGHT - CAPTION_HEIGHT

# Brand colors, matching the panes (tokens.css palette).
NAVY = RGBColor(0x14, 0x21, 0x3D)
MINT = RGBColor(0x2E, 0xC4, 0xB6)

TITLE_AND_CONTENT = 1
TITLE_ONLY = 5
PICTURE_WITH_CAPTION = 8

OUT_PATH = Path(__file__).parent / "pls,fix Demo Deck.pptx"


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def set_box(shape, left, top, width, height) -> None:
    shape.left, shape.top, shape.width, shape.height = left, top, width, height


def add_title(slide, text: str) -> None:
    """Put the title in the margin band (0..TITLE_BAND_LIMIT) above the
    content area and rule it off in mint, so the pane's free-space scan
    reads CONTENT_TOP as free rather than blocked by the title's own gap
    buffer. Centered on every slide: a layout's own alignment (Picture with
    Caption defaults to left) is not trusted."""
    title = slide.shapes.title
    title.text = text
    set_box(title, CONTENT_LEFT, Pt(0), CONTENT_WIDTH, TITLE_HEIGHT)
    paragraph = title.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    font = paragraph.font
    font.size = Pt(16)
    font.bold = True
    font.color.rgb = NAVY

    rule_top = TITLE_BAND_LIMIT - RULE_HEIGHT
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, CONTENT_LEFT, rule_top, CONTENT_WIDTH, RULE_HEIGHT)
    rule.fill.solid()
    rule.fill.fore_color.rgb = MINT
    rule.line.fill.background()
    rule.shadow.inherit = False


def add_slide(prs: Presentation, layout: int, title_text: str):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    add_title(slide, title_text)
    return slide


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def set_number_bullet(paragraph) -> None:
    """Real auto-numbering (1., 2., ...) instead of the master's inherited
    bullet char, reached via the paragraph's own oxml element - python-pptx
    has no public API for bullet formatting."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.remove_all("a:buNone", "a:buAutoNum", "a:buChar")
    pPr.insert_element_before(
        parse_xml(f'<a:buFont {nsdecls("a")} typeface="+mj-lt"/>'),
        "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )
    pPr.insert_element_before(
        parse_xml(f'<a:buAutoNum {nsdecls("a")} type="arabicPeriod"/>'),
        "a:buNone", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )


def set_no_bullet(paragraph) -> None:
    """Suppress the master's inherited bullet on a plain paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.remove_all("a:buNone", "a:buAutoNum", "a:buChar")
    pPr.insert_element_before(
        parse_xml(f'<a:buNone {nsdecls("a")}/>'),
        "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )


def build_slide_1(prs: Presentation) -> None:
    """The four PowerPoint tasks as a numbered list, plus how to run them."""
    slide = add_slide(prs, TITLE_AND_CONTENT, "pls,fix demo deck")
    body = slide.placeholders[1]
    set_box(body, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
    tf = body.text_frame
    tf.word_wrap = True

    tasks = [
        "P&L table → Slide 2, Where = Whole slide",
        "Revenue chart → Slide 3, Where = Left half",
        "Segment pie → Slide 3, Where = Right half",
        "Picture from the Data sheet → Slide 4: select the empty "
        "placeholder, Where = Selected shape",
    ]
    tf.paragraphs[0].text = tasks[0]
    tf.paragraphs[0].font.size = Pt(18)
    set_number_bullet(tf.paragraphs[0])
    for task in tasks[1:]:
        para = tf.add_paragraph()
        para.text = task
        para.font.size = Pt(18)
        set_number_bullet(para)

    commands = tf.add_paragraph()
    commands.space_before = Pt(20)
    set_no_bullet(commands)
    label = commands.add_run()
    label.text, label.font.bold, label.font.size = "Commands: ", True, Pt(14)
    rest = commands.add_run()
    rest.text = (
        "open the pane from the pls,fix ribbon tab, Inbox tab, the Slide and "
        "Where pickers above the list, Insert, Update all."
    )
    rest.font.size = Pt(14)
    set_notes(slide, "This deck demos the PowerPoint Slide and Where pickers across the four tasks above.")


def build_slide_2(prs: Presentation) -> None:
    """One empty content placeholder: free space for the pane to place into."""
    slide = add_slide(prs, TITLE_AND_CONTENT, "Choose the slide")
    body = slide.placeholders[1]
    set_box(body, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
    # No text written: an empty layout placeholder reads as free space, the
    # same rule src/ppt/placement.ts applies when it scans a slide's shapes.
    set_notes(slide, "Choose the slide: pick this slide, then a spot, in the Slide picker.")


def add_half_diagram(slide, box, caption_text: str) -> None:
    left, top, width, height = box
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.background()
    rect.line.color.rgb = MINT
    rect.line.width = Pt(2)
    rect.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    rect.shadow.inherit = False

    caption = slide.shapes.add_textbox(left, top + height + Pt(4), width, Pt(24))
    para = caption.text_frame.paragraphs[0]
    para.text = caption_text
    para.alignment = PP_ALIGN.CENTER
    para.font.size = Pt(12)
    para.font.color.rgb = NAVY


def build_slide_3(prs: Presentation) -> None:
    """The pane's left and right halves, drawn to their exact geometry."""
    slide = add_slide(prs, TITLE_ONLY, "Choose the spot")
    add_half_diagram(slide, LEFT_HALF, "Revenue chart: Where = Left half")
    add_half_diagram(slide, RIGHT_HALF, "Segment pie: Where = Right half")
    set_notes(slide, "Choose the spot: Where picks a half, a quarter, the whole slide or the selected shape.")


def build_slide_4(prs: Presentation) -> None:
    """An empty picture placeholder to select and insert into."""
    slide = add_slide(prs, PICTURE_WITH_CAPTION, "Into a placeholder")
    picture = slide.placeholders[1]
    set_box(picture, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, PICTURE_HEIGHT)
    # Left empty: no insert_picture() call, so it stays the "click to add
    # picture" placeholder the pane's occupied-box scan treats as free space.

    caption = slide.placeholders[2]
    set_box(caption, CONTENT_LEFT, CAPTION_TOP, CONTENT_WIDTH, CAPTION_HEIGHT)
    para = caption.text_frame.paragraphs[0]
    para.text = "Select the placeholder, Where = Selected shape, Insert."
    para.font.size = Pt(14)
    para.font.color.rgb = NAVY
    set_notes(slide, "Into a placeholder: select the empty picture placeholder, Where = Selected shape, Insert.")


def set_core_properties(prs: Presentation) -> None:
    """Overwrite the default template's stale docProps/core.xml: its author,
    description and 2013 timestamps are python-pptx's own template, not this
    project or this build."""
    props = prs.core_properties
    props.title = "pls,fix Demo Deck"
    props.author = "pls,fix"
    props.last_modified_by = "pls,fix"
    props.comments = ""
    now = datetime.now(timezone.utc).replace(microsecond=0)
    props.created = now
    props.modified = now


def drop_template_baggage(prs: Presentation) -> None:
    """Remove docProps/thumbnail.jpeg (a stale 4:3 preview) and
    ppt/printerSettings/printerSettings1.bin (Windows print-driver settings),
    both dead weight from python-pptx's template: drop the one relationship
    each has, so OpcPackage.iter_parts()'s reachability walk excludes the
    part - and its [Content_Types].xml entry - from save() on its own. No
    zip surgery; proved clean by unzip -t and a reload round-trip."""
    package = prs.part.package
    pkg_rels = package._rels  # OpcPackage has no public alias for this
    thumbnail_rid = next(rid for rid in pkg_rels if pkg_rels[rid].reltype == RT.THUMBNAIL)
    package.drop_rel(thumbnail_rid)

    pres_rels = prs.part.rels
    printer_rid = next(rid for rid in pres_rels if pres_rels[rid].reltype == RT.PRINTER_SETTINGS)
    prs.part.drop_rel(printer_rid)


def build() -> Presentation:
    prs = new_presentation()
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    set_core_properties(prs)
    drop_template_baggage(prs)
    return prs


def main() -> None:
    prs = build()
    prs.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
